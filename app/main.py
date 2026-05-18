"""Main user-facing blueprint: landing page, generation, profile."""
from __future__ import annotations

import hmac
import io
import json
import logging
import os
import traceback
from datetime import datetime

from flask import (Blueprint, current_app, flash, redirect, render_template,
                   request, send_file, session, url_for, jsonify)
from flask_login import current_user, login_required

from . import db as dbm
from . import xray_service as xs


main_bp = Blueprint("main", __name__)
logger = logging.getLogger(__name__)


def _require_agent_token() -> tuple[bool, str]:
    """Validate X-Agent-Token header against AGENT_API_TOKEN env var."""
    configured = (os.environ.get("AGENT_API_TOKEN") or "").strip()
    if not configured:
        return False, "AGENT_API_TOKEN is not configured on the server."

    provided = (request.headers.get("X-Agent-Token") or "").strip()
    if not provided:
        return False, "Missing X-Agent-Token header."

    if not hmac.compare_digest(provided, configured):
        return False, "Invalid agent token."

    return True, ""


def _is_admin_without_enduser_access() -> bool:
    return bool(current_user.is_admin and (current_user.username or "").lower() == "admin")


def _build_site_summary(rows) -> list[dict]:
    summary = {}
    for r in rows or []:
        site_name = (r["Site_Name"] or "").strip() if "Site_Name" in r.keys() else ""
        site_name = site_name or "Unknown"
        bucket = summary.setdefault(site_name, {"jira_ids": set(), "test_ids": set()})
        bucket["jira_ids"].add((r["Jira_id"] or "").strip())
        bucket["test_ids"].add((r["Test_id"] or "").strip())

    out = []
    for site_name in sorted(summary.keys()):
        bucket = summary[site_name]
        out.append(
            {
                "Site_Name": site_name,
                "Jira_Count": len([x for x in bucket["jira_ids"] if x]),
                "Test_Count": len([x for x in bucket["test_ids"] if x]),
            }
        )
    return out


@main_bp.route("/")
def root():
    return redirect(url_for("auth.login"))


@main_bp.route("/landing")
@login_required
def landing():
    if _is_admin_without_enduser_access():
        return redirect(url_for("admin.dashboard"))
    db_path = current_app.config["DATABASE"]
    xray_rows = dbm.list_xray_tests(db_path, created_by=current_user.username)
    if not xray_rows:
        # Fallback: avoid blank grid if historical Created_By values differ unexpectedly.
        xray_rows = dbm.list_xray_tests(db_path)

    xray_site_rows = _build_site_summary(xray_rows)

    export_rows = dbm.list_exports(db_path, created_by=current_user.username)
    return render_template(
        "landing.html",
        xray_site_rows=xray_site_rows,
        export_rows=export_rows,
        has_pat=bool(current_user.jira_pat),
    )


@main_bp.route("/api/imported-tests/site-summary")
@login_required
def api_imported_tests_site_summary():
    if _is_admin_without_enduser_access():
        return jsonify({"rows": []})

    db_path = current_app.config["DATABASE"]
    rows = dbm.list_xray_tests(db_path, created_by=current_user.username)
    if not rows:
        rows = dbm.list_xray_tests(db_path)

    summary_rows = _build_site_summary(rows)
    return jsonify({"rows": summary_rows})


@main_bp.route("/api/agent/health", methods=["GET"])
def api_agent_health():
    ok, msg = _require_agent_token()
    if not ok:
        return jsonify({"ok": False, "error": msg}), 401
    return jsonify({"ok": True, "service": "jira-test-generator"})


@main_bp.route("/api/agent/generate", methods=["POST"])
def api_agent_generate():
    """Machine-callable endpoint for agents (no browser login required)."""
    ok, msg = _require_agent_token()
    if not ok:
        return jsonify({"ok": False, "error": msg}), 401

    payload = request.get_json(silent=True) or {}
    story_key = (payload.get("story_key") or "").strip()
    jira_pat = (payload.get("jira_pat") or "").strip()
    action = (payload.get("action") or "preview").strip().lower()
    custom_prompt = (payload.get("custom_prompt") or "").strip()
    created_by = (payload.get("created_by") or "agent").strip() or "agent"
    include_testcases = bool(payload.get("include_testcases", False))

    if not story_key:
        return jsonify({"ok": False, "error": "story_key is required."}), 400
    if not jira_pat:
        return jsonify({"ok": False, "error": "jira_pat is required."}), 400
    if action not in ("preview", "excel", "jira"):
        return jsonify({"ok": False, "error": "action must be one of: preview, excel, jira."}), 400

    try:
        story_context = xs.build_story_context(jira_pat, story_key)
        prompt = custom_prompt or xs.DEFAULT_PROMPT
        testcases = xs.generate_testcases(prompt, story_context)
    except Exception as e:
        logger.exception("Agent generation failed")
        return jsonify({"ok": False, "error": f"Generation failed: {e}"}), 500

    db_path = current_app.config["DATABASE"]
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports")

    response = {
        "ok": True,
        "story_key": story_key,
        "action": action,
        "testcase_count": len(testcases),
    }
    if include_testcases or action == "preview":
        response["testcases"] = testcases

    if action == "preview":
        return jsonify(response)

    if action == "excel":
        try:
            out_file = xs.export_testcases_to_excel(story_key, testcases, out_dir)
            export_id = dbm.insert_export(
                db_path,
                jira_id=story_key,
                test_name=os.path.basename(out_file),
                count=len(testcases),
                created_by=created_by,
            )
            response["export_id"] = export_id
            response["excel_filename"] = os.path.basename(out_file)
            response["excel_path"] = out_file
            return jsonify(response)
        except Exception as e:
            logger.exception("Agent excel export failed")
            return jsonify({"ok": False, "error": f"Excel export failed: {e}"}), 500

    try:
        issuetype_name = xs.detect_test_issuetype_name(jira_pat)
        steps_field_id = xs.detect_steps_field_id(jira_pat)
        tt_field_id, tt_payload = xs.detect_test_type_field(jira_pat)
    except Exception as e:
        logger.exception("Agent field detection failed")
        return jsonify({"ok": False, "error": f"Field detection failed: {e}"}), 500

    project_key = story_key.split("-")[0]
    created = []
    errors = []
    for tc in testcases:
        try:
            test_key = xs.create_xray_test(
                jira_pat,
                story_key=story_key,
                project_key=project_key,
                issuetype_name=issuetype_name,
                steps_field_id=steps_field_id,
                test_type_field_id=tt_field_id,
                test_type_payload=tt_payload,
                tc=tc,
            )
            xs.link_test_to_story(jira_pat, test_key, story_key)
            steps_text, expected_text = xs.steps_as_text(tc)
            dbm.insert_xray_test(
                db_path,
                jira_id=story_key,
                test_id=test_key,
                description=xs.build_description(story_key, tc),
                steps=steps_text,
                expected=expected_text,
                created_by=created_by,
            )
            created.append(test_key)
        except Exception as e:
            logger.exception("Agent failed creating test for %s", tc.get("Title"))
            errors.append(f"{tc.get('Title', '?')}: {e}")

    response["created"] = created
    response["errors"] = errors
    response["created_count"] = len(created)
    return jsonify(response)


@main_bp.route("/imported-tests/site")
@login_required
def imported_tests_site_detail():
    if _is_admin_without_enduser_access():
        return redirect(url_for("admin.dashboard"))

    site_name = (request.args.get("site") or "").strip()
    if not site_name:
        flash("Site name is required.", "error")
        return redirect(url_for("main.landing"))

    db_path = current_app.config["DATABASE"]
    rows = dbm.list_xray_tests_by_site(db_path, site_name, created_by=current_user.username)
    if not rows:
        rows = dbm.list_xray_tests_by_site(db_path, site_name)

    if not rows:
        all_rows = dbm.list_xray_tests(db_path, created_by=current_user.username)
        if not all_rows:
            all_rows = dbm.list_xray_tests(db_path)
        rows = [
            r for r in all_rows
            if ((r["Site_Name"] or "").strip() or "Unknown") == site_name
        ]

    return render_template("imported_tests_site_detail.html", site_name=site_name, rows=rows)


@main_bp.route("/exports/<int:sl_no>")
@login_required
def export_detail(sl_no: int):
    if _is_admin_without_enduser_access():
        return redirect(url_for("admin.dashboard"))
    db_path = current_app.config["DATABASE"]
    row = dbm.get_export(db_path, sl_no)
    if not row or row["Created_By"] != current_user.username:
        flash("Export not found.", "error")
        return redirect(url_for("main.landing"))

    out_dir = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports")
    file_path = os.path.join(out_dir, row["Test_name"])
    headers, rows = [], []
    error = None
    if not os.path.exists(file_path):
        error = "Excel file is no longer available on the server."
    else:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            for i, r in enumerate(it):
                values = ["" if v is None else str(v) for v in r]
                if i == 0:
                    headers = values
                else:
                    rows.append(values)
            wb.close()
        except Exception as e:
            error = f"Failed to read Excel: {e}"

    return render_template("export_detail.html", export=row,
                           headers=headers, rows=rows, error=error)


@main_bp.route("/exports/<int:sl_no>/download")
@login_required
def export_download(sl_no: int):
    if _is_admin_without_enduser_access():
        return redirect(url_for("admin.dashboard"))
    db_path = current_app.config["DATABASE"]
    row = dbm.get_export(db_path, sl_no)
    if not row or row["Created_By"] != current_user.username:
        flash("Export not found.", "error")
        return redirect(url_for("main.landing"))

    out_dir = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports")
    file_path = os.path.join(out_dir, row["Test_name"])
    if not os.path.exists(file_path):
        flash("Excel file is no longer available on the server.", "error")
        return redirect(url_for("main.landing"))
    return send_file(file_path, as_attachment=True,
                     download_name=os.path.basename(file_path))


@main_bp.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    if _is_admin_without_enduser_access():
        return redirect(url_for("admin.dashboard"))
    db_path = current_app.config["DATABASE"]
    if request.method == "POST":
        new_pat = (request.form.get("jira_pat") or "").strip()
        dbm.update_jira_pat(db_path, current_user.id, new_pat or None)
        flash("Jira PAT updated.", "success")
        return redirect(url_for("main.profile"))
    return render_template("profile.html")


@main_bp.route("/generate", methods=["POST"])
@login_required
def generate():
    """Step 1: fetch story + generate testcases. Stash in session for confirmation."""
    if _is_admin_without_enduser_access():
        flash("Admins must use a non-admin account to create tests.", "error")
        return redirect(url_for("admin.dashboard"))

    story_key = (request.form.get("story_key") or "").strip()
    action = request.form.get("action")  # "excel" or "jira"
    custom_prompt = (request.form.get("custom_prompt") or "").strip()

    if not story_key:
        flash("Please enter a Jira story key.", "error")
        return redirect(url_for("main.landing"))
    if action not in ("excel", "jira"):
        flash("Please select an action.", "error")
        return redirect(url_for("main.landing"))
    if not current_user.jira_pat:
        flash("Please set your Jira PAT in your profile first.", "error")
        return redirect(url_for("main.profile"))

    try:
        story_context = xs.build_story_context(current_user.jira_pat, story_key)
        prompt = custom_prompt or xs.DEFAULT_PROMPT
        testcases = xs.generate_testcases(prompt, story_context)
    except Exception as e:
        logger.exception("Generation failed")
        flash(f"Generation failed: {e}", "error")
        return redirect(url_for("main.landing"))

    db_path = current_app.config["DATABASE"]
    pending_id = dbm.upsert_pending_generation(
        db_path,
        user_id=current_user.id,
        story_key=story_key,
        action=action,
        testcases=testcases,
    )
    session["pending_id"] = pending_id
    return render_template("confirm.html", story_key=story_key, action=action,
                           testcases=testcases)


@main_bp.route("/confirm", methods=["POST"])
@login_required
def confirm():
    decision = request.form.get("decision")  # yes | regenerate | cancel
    db_path = current_app.config["DATABASE"]
    pending_id = session.get("pending_id")
    pending = None
    if pending_id:
        pending = dbm.get_pending_generation(db_path, int(pending_id), current_user.id)
    if not pending:
        flash("No pending generation; please start again.", "error")
        return redirect(url_for("main.landing"))

    story_key = pending["story_key"]
    action = pending["action"]
    testcases = pending["testcases"]

    if decision == "cancel":
        dbm.clear_pending_generation(db_path, int(pending["id"]), current_user.id)
        session.pop("pending_id", None)
        flash("Cancelled.", "info")
        return redirect(url_for("main.landing"))

    if decision == "regenerate":
        custom = (request.form.get("custom_prompt") or "").strip()
        if not custom:
            flash("Please provide a custom prompt to regenerate.", "error")
            return render_template("confirm.html", story_key=story_key,
                                   action=action, testcases=testcases)
        try:
            ctx = xs.build_story_context(current_user.jira_pat, story_key)
            testcases = xs.generate_testcases(custom, ctx)
        except Exception as e:
            logger.exception("Regeneration failed")
            flash(f"Regeneration failed: {e}", "error")
            return redirect(url_for("main.landing"))
        pending_id = dbm.upsert_pending_generation(
            db_path,
            user_id=current_user.id,
            story_key=story_key,
            action=action,
            testcases=testcases,
        )
        session["pending_id"] = pending_id
        return render_template("confirm.html", story_key=story_key,
                               action=action, testcases=testcases)

    # decision == yes -> execute the action
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                           "exports")

    if action == "excel":
        try:
            out_file = xs.export_testcases_to_excel(story_key, testcases, out_dir)
        except Exception as e:
            logger.exception("Excel export failed")
            flash(f"Excel export failed: {e}", "error")
            return redirect(url_for("main.landing"))

        export_id = dbm.insert_export(
            db_path,
            jira_id=story_key,
            test_name=os.path.basename(out_file),
            count=len(testcases),
            created_by=current_user.username,
        )
        dbm.clear_pending_generation(db_path, int(pending["id"]), current_user.id)
        session.pop("pending_id", None)
        return render_template(
            "download_success.html",
            filename=os.path.basename(out_file),
            download_url=url_for("main.export_download", sl_no=export_id),
            home_url=url_for("main.landing"),
        )

    # action == jira
    try:
        issuetype_name = xs.detect_test_issuetype_name(current_user.jira_pat)
        steps_field_id = xs.detect_steps_field_id(current_user.jira_pat)
        tt_field_id, tt_payload = xs.detect_test_type_field(current_user.jira_pat)
    except Exception as e:
        logger.exception("Field detection failed")
        flash(f"Field detection failed: {e}", "error")
        return redirect(url_for("main.landing"))

    project_key = story_key.split("-")[0]
    created = []
    errors = []
    for tc in testcases:
        try:
            test_key = xs.create_xray_test(
                current_user.jira_pat,
                story_key=story_key,
                project_key=project_key,
                issuetype_name=issuetype_name,
                steps_field_id=steps_field_id,
                test_type_field_id=tt_field_id,
                test_type_payload=tt_payload,
                tc=tc,
            )
            xs.link_test_to_story(current_user.jira_pat, test_key, story_key)
            steps_text, expected_text = xs.steps_as_text(tc)
            dbm.insert_xray_test(
                db_path,
                jira_id=story_key,
                test_id=test_key,
                description=xs.build_description(story_key, tc),
                steps=steps_text,
                expected=expected_text,
                created_by=current_user.username,
            )
            created.append(test_key)
        except Exception as e:
            logger.exception("Failed creating test for %s", tc.get("Title"))
            errors.append(f"{tc.get('Title','?')}: {e}")

    # Also export an Excel and record export history
    try:
        out_file = xs.export_testcases_to_excel(story_key, testcases, out_dir)
        dbm.insert_export(
            db_path,
            jira_id=story_key,
            test_name=os.path.basename(out_file),
            count=len(testcases),
            created_by=current_user.username,
        )
    except Exception:
        out_file = None

    dbm.clear_pending_generation(db_path, int(pending["id"]), current_user.id)
    session.pop("pending_id", None)
    return render_template(
        "result.html",
        story_key=story_key,
        created=created,
        errors=errors,
        excel_filename=os.path.basename(out_file) if out_file else None,
    )
