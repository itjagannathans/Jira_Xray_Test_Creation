"""Admin blueprint: dashboard, user management, history views, Excel export."""
from __future__ import annotations

import io
from datetime import datetime
from functools import wraps

from flask import (Blueprint, abort, current_app, flash, redirect,
                   render_template, request, send_file, url_for)
from flask_login import current_user, login_required
from openpyxl import Workbook
from werkzeug.security import generate_password_hash

from . import db as dbm


admin_bp = Blueprint("admin", __name__)


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapper(*args, **kwargs):
        if not current_user.is_admin:
            abort(403)
        return view(*args, **kwargs)
    return wrapper


@admin_bp.route("/")
@admin_required
def dashboard():
    db_path = current_app.config["DATABASE"]
    users = dbm.list_users(db_path)
    xray_count = len(dbm.list_xray_tests(db_path))
    export_count = len(dbm.list_exports(db_path))
    return render_template(
        "admin/dashboard.html",
        users=users,
        xray_count=xray_count,
        export_count=export_count,
    )


@admin_bp.route("/users")
@admin_required
def users():
    db_path = current_app.config["DATABASE"]
    return render_template("admin/users.html", users=dbm.list_users(db_path))


@admin_bp.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_user(user_id: int):
    db_path = current_app.config["DATABASE"]
    row = dbm.get_user_by_id(db_path, user_id)
    if not row:
        abort(404)

    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        jira_pat = (request.form.get("jira_pat") or "").strip() or None
        is_admin = bool(request.form.get("is_admin"))
        new_password = request.form.get("new_password") or ""

        if not username:
            flash("Username cannot be empty.", "error")
            return render_template("admin/edit_user.html", user=row)

        dbm.update_user(db_path, user_id, username=username,
                        jira_pat=jira_pat, is_admin=is_admin)
        if new_password:
            if len(new_password) < 8:
                flash("Password must be at least 8 characters.", "error")
                return render_template("admin/edit_user.html", user=row)
            with dbm.connect(db_path) as conn:
                conn.execute("UPDATE Tb_user SET password_hash = ? WHERE id = ?",
                             (generate_password_hash(new_password), user_id))
        flash("User updated.", "success")
        return redirect(url_for("admin.users"))

    return render_template("admin/edit_user.html", user=row)


@admin_bp.route("/users/<int:user_id>/toggle-active", methods=["POST"])
@admin_required
def toggle_active(user_id: int):
    db_path = current_app.config["DATABASE"]
    row = dbm.get_user_by_id(db_path, user_id)
    if not row:
        abort(404)
    if row["username"] == current_user.username:
        flash("You cannot deactivate your own account.", "error")
        return redirect(url_for("admin.users"))
    dbm.set_user_active(db_path, user_id, not bool(row["is_active"]))
    flash("User status updated.", "success")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/<int:user_id>/delete", methods=["POST"])
@admin_required
def delete_user(user_id: int):
    db_path = current_app.config["DATABASE"]
    row = dbm.get_user_by_id(db_path, user_id)
    if not row:
        abort(404)
    if row["username"] == current_user.username:
        flash("You cannot delete your own account.", "error")
        return redirect(url_for("admin.users"))
    dbm.delete_user(db_path, user_id)
    flash("User deleted.", "success")
    return redirect(url_for("admin.users"))


@admin_bp.route("/xray-history")
@admin_required
def xray_history():
    db_path = current_app.config["DATABASE"]
    return render_template("admin/xray_history.html",
                           rows=dbm.list_xray_tests(db_path))


@admin_bp.route("/export-history")
@admin_required
def export_history():
    db_path = current_app.config["DATABASE"]
    return render_template("admin/export_history.html",
                           rows=dbm.list_exports(db_path))


@admin_bp.route("/sitewise-report")
@admin_required
def sitewise_report():
    db_path = current_app.config["DATABASE"]
    rows = dbm.list_sitewise_report_summary(db_path)
    return render_template("admin/sitewise_report.html", rows=rows)


@admin_bp.route("/export-history/<int:sl_no>")
@admin_required
def export_history_detail(sl_no: int):
    import os as _os
    from openpyxl import load_workbook
    db_path = current_app.config["DATABASE"]
    row = dbm.get_export(db_path, sl_no)
    if not row:
        flash("Export not found.", "error")
        return redirect(url_for("admin.export_history"))

    out_dir = _os.path.join(
        _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "exports")
    file_path = _os.path.join(out_dir, row["Test_name"])
    headers, data_rows = [], []
    error = None
    if not _os.path.exists(file_path):
        error = "Excel file is no longer available on the server."
    else:
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                values = ["" if v is None else str(v) for v in r]
                if i == 0:
                    headers = values
                else:
                    data_rows.append(values)
            wb.close()
        except Exception as e:
            error = f"Failed to read Excel: {e}"

    return render_template("admin/export_detail.html", export=row,
                           headers=headers, rows=data_rows, error=error)


def _fmt_dt(v):
    return str(v).replace("T", " ") if v else ""


def _build_userwise_rows(db_path: str):
    """Aggregate per-user counts of stories/tests for Excel and Xray."""
    query = """
        SELECT
            u.username AS username,
            COALESCE(NULLIF(u.full_name, ''), u.username) AS display_name,
            COALESCE(ex.stories_excel, 0) AS stories_excel,
            COALESCE(ex.tests_excel, 0) AS tests_excel,
            COALESCE(xr.stories_xray, 0) AS stories_xray,
            COALESCE(xr.tests_xray, 0) AS tests_xray,
            CASE
                WHEN ex.last_export IS NULL THEN xr.last_xray
                WHEN xr.last_xray IS NULL THEN ex.last_export
                WHEN ex.last_export >= xr.last_xray THEN ex.last_export
                ELSE xr.last_xray
            END AS last_activity
        FROM Tb_user u
        LEFT JOIN (
            SELECT
                Created_By,
                COUNT(DISTINCT Jira_id) AS stories_excel,
                COALESCE(SUM(NoOfTest_Created), 0) AS tests_excel,
                MAX(Created_date) AS last_export
            FROM Tb_Testcase_Export
            GROUP BY Created_By
        ) ex ON ex.Created_By = u.username
        LEFT JOIN (
            SELECT
                Created_By,
                COUNT(DISTINCT Jira_id) AS stories_xray,
                COUNT(*) AS tests_xray,
                MAX(Created_date) AS last_xray
            FROM Tb_Xray_Testcase_Created
            GROUP BY Created_By
        ) xr ON xr.Created_By = u.username
        WHERE LOWER(COALESCE(u.username, '')) <> 'admin'
        ORDER BY LOWER(COALESCE(NULLIF(u.full_name, ''), u.username)) ASC
    """

    rows = []
    with dbm.connect(db_path) as conn:
        result = conn.execute(query).fetchall()

    for sl, r in enumerate(result, start=1):
        rows.append({
            "sl_no": sl,
            "username": r["username"],
            "display_name": r["display_name"],
            "last_activity": _fmt_dt(r["last_activity"]),
            "excel_stories": int(r["stories_excel"] or 0),
            "excel_tests": int(r["tests_excel"] or 0),
            "xray_stories": int(r["stories_xray"] or 0),
            "xray_tests": int(r["tests_xray"] or 0),
        })
    return rows


@admin_bp.route("/userwise-history")
@admin_required
def userwise_history():
    db_path = current_app.config["DATABASE"]
    return render_template("admin/userwise_history.html",
                           rows=_build_userwise_rows(db_path))


@admin_bp.route("/userwise-history/export")
@admin_required
def userwise_history_export():
    db_path = current_app.config["DATABASE"]
    rows = _build_userwise_rows(db_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Userwise History"
    ws.append(["Sl No", "User Name", "Email", "Last Activity Date",
               "No of Stories - Excel", "No of Test Generated - Excel",
               "No of Stories - Xray", "No of Test Generated - Xray"])
    for r in rows:
        ws.append([r["sl_no"], r["display_name"], r["username"], r["last_activity"],
                   r["excel_stories"], r["excel_tests"],
                   r["xray_stories"], r["xray_tests"]])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Userwise_History_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@admin_bp.route("/export-activity")
@admin_required
def export_activity():
    db_path = current_app.config["DATABASE"]
    users = dbm.list_users(db_path)
    xrays = dbm.list_xray_tests(db_path)
    exports = dbm.list_exports(db_path)

    name_lookup = {u["username"]: (u["full_name"] or u["username"]) for u in users}

    wb = Workbook()

    ws_users = wb.active
    ws_users.title = "Users"
    ws_users.append(["Id", "Name", "Email", "Is Admin", "Is Active", "Created Date"])
    for u in users:
        ws_users.append([u["id"], (u["full_name"] or u["username"]), u["username"],
                         "Yes" if u["is_admin"] else "No",
                         "Yes" if u["is_active"] else "No",
                         _fmt_dt(u["created_date"])])

    ws_x = wb.create_sheet("Xray Testcases Created")
    ws_x.append(["SL_No", "Jira_id", "Test_id", "Test_Description",
                 "Test_Steps", "Test_expected", "Created By (Name)", "Created By (Email)", "Created_date"])
    for r in xrays:
        ws_x.append([r["SL_No"], r["Jira_id"], r["Test_id"],
                     r["Test_Description"], r["Test_Steps"], r["Test_expected"],
                     name_lookup.get(r["Created_By"], r["Created_By"]),
                     r["Created_By"], _fmt_dt(r["Created_date"])])

    ws_e = wb.create_sheet("Manual Test Exports")
    ws_e.append(["Sl_No", "Jira_id", "Test_name", "NoOfTest_Created",
                 "Created By (Name)", "Created By (Email)", "Created_date"])
    for r in exports:
        ws_e.append([r["Sl_No"], r["Jira_id"], r["Test_name"],
                     r["NoOfTest_Created"],
                     name_lookup.get(r["Created_By"], r["Created_By"]),
                     r["Created_By"], _fmt_dt(r["Created_date"])])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"All_User_Activity_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
